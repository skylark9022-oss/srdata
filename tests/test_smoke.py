"""패키지가 임포트되고 기본 기능이 도는지 확인하는 최소 테스트."""

import pandas as pd

from srdata import config
from srdata.excel import save_sheets
from srdata.excel.writer import _safe_sheet_name


def test_paths_are_under_root():
    assert config.RAW_DIR.is_relative_to(config.ROOT_DIR)
    assert config.PROCESSED_DIR.is_relative_to(config.ROOT_DIR)


def test_save_sheets_roundtrip(tmp_path):
    df = pd.DataFrame({"이름": ["가", "나"], "값": [1, 2]})
    # "표/1" 과 "표:1" 은 금지문자 치환 후 둘 다 "표_1" 이 되므로 두 번째는 "_2" 가 붙어야 한다.
    out = save_sheets(tmp_path / "t.xlsx", {"표/1": df, "표:1": df})
    back = pd.read_excel(out, sheet_name=None)
    assert list(back) == ["표_1", "표_1_2"]
    assert back["표_1"].equals(df)


def test_sheet_name_dedup_and_length():
    used: set[str] = set()
    a = _safe_sheet_name("x" * 40, used)
    b = _safe_sheet_name("x" * 40, used)
    assert len(a) == 31 and len(b) == 31 and a != b


def test_numeric_text_keeps_decimals_and_codes(tmp_path):
    import openpyxl
    df = pd.DataFrame({"v": ["1.00", "<0.43", "0", "0.50"], "code": ["0012", "102-04", "7", "0.5"]})
    out = save_sheets(tmp_path / "n.xlsx", {"d": df}, numeric_text={"d"}, keep_leading_zeros=True)
    ws = openpyxl.load_workbook(out)["d"]
    assert [(c.value, c.number_format) for c in ws["A"][1:]] == [(1.0, "0.00"), ("<0.43", "General"), (0, "General"), (0.5, "0.00")]
    assert [c.value for c in ws["B"][1:]] == ["0012", "102-04", 7, 0.5]
