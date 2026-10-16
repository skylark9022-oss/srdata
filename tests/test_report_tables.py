"""report_tables 의 순수 함수(병합 셀 해석, 머리글 합치기) 테스트."""

from srdata.pdf.report_tables import (
    _marker_diff,
    coerce_number,
    flatten_headers,
    header_row_count,
    resolve_grid,
)


def _box(x0, top, x1, bottom):
    return (x0, top, x1, bottom)


# 3열 표. 머리글 2행: 열0 '해역' 은 세로 병합, 열1~2 는 '농도' 아래에 '2025년','과거' 가 있다.
# 본문 2행: 열0 '동해' 가 세로 병합. 마지막 행은 '평균' 이 열0~1 가로 병합.
CELLS = [
    [_box(0, 0, 10, 20), _box(10, 0, 30, 10), None],
    [None, _box(10, 10, 20, 20), _box(20, 10, 30, 20)],
    [_box(0, 20, 10, 40), _box(10, 20, 20, 30), _box(20, 20, 30, 30)],
    [None, _box(10, 30, 20, 40), _box(20, 30, 30, 40)],
    [_box(0, 40, 20, 50), None, _box(20, 40, 30, 50)],
]
TEXTS = [
    ["해역", "농도", None],
    [None, "2025년", "과거"],
    ["동해", "1.0", "0.5~1.5"],
    [None, "2.0", ""],
    ["평균", None, "1.5"],
]


def test_resolve_grid_marks_vertical_and_horizontal_merges():
    grid = resolve_grid(CELLS, TEXTS)
    assert grid[0][2] == ("농도", "hmerge")
    assert grid[1][0] == ("해역", "vmerge")
    assert grid[3][0] == ("동해", "vmerge")
    assert grid[3][2] == ("", "cell")  # 진짜 빈 셀은 cell 로 남는다
    assert grid[4][1] == ("평균", "hmerge")


def test_header_row_count_and_flatten():
    assert header_row_count(CELLS) == 2
    grid = resolve_grid(CELLS, TEXTS)
    assert flatten_headers(grid, 2) == ["해역", "농도 / 2025년", "농도 / 과거"]


def test_flatten_headers_makes_names_unique():
    grid = [[("a", "cell"), ("a", "hmerge"), ("", "cell")]]
    assert flatten_headers(grid, 1) == ["a", "a_1", "col3"]


def test_marker_diff_and_numbers():
    assert _marker_diff("<0.0011)", "<0.001") == "1)"
    assert _marker_diff("137Cs1)", "137Cs") == "1)"
    assert coerce_number("12.2") == 12.2
    assert coerce_number("0") == 0
    assert coerce_number("102-04") == "102-04"
    assert coerce_number("2025.05.29") == "2025.05.29"


def test_numeric_strings_to_numbers_keeps_decimals(tmp_path):
    import pandas as pd
    from openpyxl import load_workbook

    from srdata.excel import save_sheets

    df = pd.DataFrame({"a": ["1.60", "0", "102-04", "<0.001"], "b": ["2025.05.29", "12", "-3.5", ""]})
    path = save_sheets(tmp_path / "t.xlsx", {"s": df}, numeric_text=True)
    ws = load_workbook(path)["s"]
    assert ws["A2"].value == 1.6 and ws["A2"].number_format == "0.00"
    assert ws["A3"].value == 0
    assert ws["A4"].value == "102-04" and ws["A5"].value == "<0.001"
    assert ws["B2"].value == "2025.05.29" and ws["B3"].value == 12 and ws["B4"].value == -3.5
